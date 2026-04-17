from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


class DocumentParser:
    """Document parser with Docling-first strategy and safe fallbacks."""

    def __init__(self) -> None:
        self._docling_converter = None
        self._docling_chunker = None
        self._try_init_docling()

    def _try_init_docling(self) -> None:
        try:
            from docling.document_converter import DocumentConverter
            from docling_core.transforms.chunker import HierarchicalChunker

            self._docling_converter = DocumentConverter()
            self._docling_chunker = HierarchicalChunker()
        except Exception:
            self._docling_converter = None
            self._docling_chunker = None

    def parse(self, file_path: str | Path) -> dict:
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix in {".txt", ".md", ".csv", ".html", ".htm"}:
            markdown = self._fallback_extract_text(path)
            chunks = self._split_text(markdown)
            return {
                "file_name": path.name,
                "markdown": markdown,
                "chunks": chunks,
                "word_count": self._count_words(markdown),
                "metadata": {
                    "source": str(path),
                    "format": suffix,
                    "parser": "fallback",
                },
            }

        if self._docling_converter is not None:
            try:
                result = self._docling_converter.convert(str(path))
                doc = result.document
                chunks = []
                if self._docling_chunker is not None:
                    chunks = [chunk.text for chunk in self._docling_chunker.chunk(doc)]
                markdown = doc.export_to_markdown()
                return {
                    "file_name": path.name,
                    "markdown": markdown,
                    "chunks": chunks or self._split_text(markdown),
                    "word_count": self._count_words(markdown),
                    "metadata": {
                        "source": str(path),
                        "format": suffix,
                        "parser": "docling",
                    },
                }
            except Exception:
                pass

        markdown = self._fallback_extract_text(path)
        chunks = self._split_text(markdown)
        return {
            "file_name": path.name,
            "markdown": markdown,
            "chunks": chunks,
            "word_count": self._count_words(markdown),
            "metadata": {
                "source": str(path),
                "format": suffix,
                "parser": "fallback",
            },
        }

    def parse_batch(self, file_paths: list[str | Path]) -> list[dict]:
        return [self.parse(path) for path in file_paths]

    @staticmethod
    def _count_words(text: str) -> int:
        import re
        if not text:
            return 0
        chinese = len(re.findall(r'[\u4e00-\u9fff]', text))
        english = len(re.findall(r'[a-zA-Z]+', text))
        return chinese + english

    @staticmethod
    def _split_text(text: str, chunk_size: int = 500, overlap: int = 100) -> list[str]:
        content = (text or "").strip()
        if not content:
            return []
        chunks = []
        start = 0
        while start < len(content):
            end = start + chunk_size
            chunks.append(content[start:end])
            start = end - overlap
        return chunks

    def _fallback_extract_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".md", ".txt"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".docx":
            from docx import Document as DocxDocument

            doc = DocxDocument(path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            table_rows: list[str] = []
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        table_rows.append(" | ".join(cells))
            return "\n".join(paragraphs + table_rows)
        if suffix in {".xlsx", ".xls"}:
            wb = load_workbook(path, read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in wb.worksheets:
                lines.append(f"# Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    row_values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_values:
                        lines.append(" | ".join(row_values))
            wb.close()
            return "\n".join(lines)
        if suffix == ".csv":
            import csv
            text = path.read_text(encoding="utf-8", errors="ignore")
            lines: list[str] = []
            for row in csv.reader(text.splitlines()):
                if any(cell.strip() for cell in row):
                    lines.append(" | ".join(cell.strip() for cell in row))
            return "\n".join(lines)
        if suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup
            html = path.read_text(encoding="utf-8", errors="ignore")
            soup = BeautifulSoup(html, "html.parser")
            return soup.get_text(separator="\n", strip=True)
        return path.read_text(encoding="utf-8", errors="ignore")