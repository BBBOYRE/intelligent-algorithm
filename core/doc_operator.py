from __future__ import annotations
import json
import re
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from langchain_core.prompts import ChatPromptTemplate

from core.document_parser import DocumentParser
from core.llm_factory import create_llm
from utils.file_utils import build_output_path


class DocOperator:
    """Execute natural language instructions on documents."""

    def __init__(self) -> None:
        self.llm = create_llm()
        self.parser = DocumentParser()

    def execute(self, file_path: str | Path, instruction: str) -> dict[str, Any]:
        path = Path(file_path)
        suffix = path.suffix.lower()

        # docx 文件：支持格式修改 + 内容编辑，生成新文件可下载
        if suffix == ".docx":
            return self._execute_docx(path, instruction)

        # 其他格式：纯文本处理，返回结果文本
        return self._execute_text(path, instruction)

    def _execute_docx(self, file_path: Path, instruction: str) -> dict[str, Any]:
        """处理 docx 文件：先让 LLM 生成操作指令 JSON，再执行格式/内容修改"""
        parsed = self.parser.parse(file_path)
        text = parsed["markdown"][:8000]

        # 第一步：让 LLM 分析指令，生成结构化操作列表
        plan_prompt = ChatPromptTemplate.from_messages([
            (
                "system",
                """你是文档操作规划专家。用户会给你一段文档内容和一条操作指令。
请分析指令，返回一个 JSON 对象，包含两部分：

1. "operations": 一个操作数组，每个操作是一个对象，支持以下类型：
   - {{"type": "replace", "old": "原文本", "new": "新文本"}} — 替换内容
   - {{"type": "delete", "target": "要删除的文本片段"}} — 删除内容
   - {{"type": "format", "target": "目标文本片段", "bold": true/false, "italic": true/false, "font_size": 数字, "align": "left/center/right"}} — 格式调整（字段可选）
   - {{"type": "insert", "after": "在此文本之后插入", "content": "要插入的内容"}} — 插入内容

2. "summary": 一句话描述你做了什么操作

严格规则：
- 只输出 JSON，不要输出任何其他文字。
- **target/old 字段的值必须是文档中实际存在的完整文本片段，不能截断或拆分。**
- 如果指令是"删除人名"，必须识别出文档中每一个完整的姓名（如"张三"、"李四"），每个姓名作为一个独立的 delete 操作，target 写完整姓名。
- **禁止按单字删除**：不能只删姓（如"周"）保留名，这会导致误删普通文字中的同姓字。
- 如果指令是提取或总结类操作（不需要修改文档），返回 {{"operations": [], "summary": "提取/总结操作", "extract_result": "提取或总结的结果文本"}}""",
            ),
            ("human", "文档内容：\n{document}\n\n操作指令：{instruction}"),
        ])

        chain = plan_prompt | self.llm
        response = chain.invoke({"document": text, "instruction": instruction})
        content = getattr(response, "content", "")
        if isinstance(content, list):
            content = "\n".join(str(item) for item in content)

        # 解析 LLM 返回的操作计划
        plan = self._parse_json_object(content)
        if plan is None:
            # 解析失败，回退到纯文本模式
            return self._execute_text(Path(file_path), instruction)

        # 如果是提取/总结操作，不修改文件
        if not plan.get("operations") and plan.get("extract_result"):
            return {
                "status": "success",
                "original_file": file_path.name,
                "instruction": instruction,
                "result": plan["extract_result"],
                "output_path": None,
            }

        # 第二步：在 docx 上执行操作
        doc = DocxDocument(file_path)
        ops_applied = 0

        for op in plan.get("operations", []):
            op_type = op.get("type", "")
            try:
                if op_type == "replace":
                    ops_applied += self._apply_replace(doc, op.get("old", ""), op.get("new", ""))
                elif op_type == "delete":
                    ops_applied += self._apply_replace(doc, op.get("target", ""), "")
                elif op_type == "format":
                    ops_applied += self._apply_format(doc, op)
                elif op_type == "insert":
                    ops_applied += self._apply_insert(doc, op.get("after", ""), op.get("content", ""))
            except Exception:
                continue

        # 保存修改后的文件
        output_path = build_output_path(f"edited_{file_path.name}")
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)

        return {
            "status": "success",
            "original_file": file_path.name,
            "instruction": instruction,
            "result": plan.get("summary", f"已执行 {ops_applied} 项操作"),
            "ops_applied": ops_applied,
            "output_path": str(output_path),
        }

    def _execute_text(self, file_path: Path, instruction: str) -> dict[str, Any]:
        """纯文本处理模式"""
        parsed = self.parser.parse(file_path)
        text = parsed["markdown"][:8000]

        prompt = ChatPromptTemplate.from_messages([
            (
                "system",
                """你是文档操作助手。用户会给你一段文档内容和一条操作指令。
请根据指令对文档进行操作，并返回结果。

操作类型包括但不限于：
- 信息提取（人名、日期、金额、关键数据等）
- 格式转换（日期格式、数字格式等）
- 内容编辑（删除、替换、重排、合并等）
- 内容总结、摘要或分析
- 排版调整建议

规则：
1. 直接返回操作后的结果，不要解释操作过程。
2. 如果是提取操作，以清晰的列表形式返回提取到的信息。
3. 如果是编辑操作，返回修改后的完整文档内容。
4. 如果是总结/分析操作，返回结构化的分析结果。
5. 回答语言必须是中文。""",
            ),
            ("human", "文档内容：\n{document}\n\n操作指令：{instruction}"),
        ])

        chain = prompt | self.llm
        response = chain.invoke({"document": text, "instruction": instruction})
        content = getattr(response, "content", "")
        if isinstance(content, list):
            content = "\n".join(str(part) for part in content)

        return {
            "status": "success",
            "original_file": file_path.name,
            "instruction": instruction,
            "result": str(content).strip() or "未生成有效结果，请重试。",
            "output_path": None,
        }

    @staticmethod
    def _apply_replace(doc: DocxDocument, old_text: str, new_text: str) -> int:
        if not old_text:
            return 0
        paragraphs = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        total_count = 0
        for para in paragraphs:
            if old_text not in para.text:
                continue
            para_count = 0
            for run in para.runs:
                if old_text in run.text:
                    run.text = run.text.replace(old_text, new_text)
                    para_count += 1
            # 如果 runs 里没匹配到（文本跨 run），整段替换
            if para_count == 0 and old_text in para.text:
                full = para.text.replace(old_text, new_text)
                for i, run in enumerate(para.runs):
                    run.text = full if i == 0 else ""
                para_count = 1
            total_count += para_count
        return total_count

    @staticmethod
    def _apply_format(doc: DocxDocument, op: dict) -> int:
        target = op.get("target", "")
        if not target:
            return 0
        paragraphs = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        total_count = 0
        for para in paragraphs:
            if target not in para.text:
                continue
            # 段落级格式
            if "align" in op:
                align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
                para.alignment = align_map.get(op["align"], WD_ALIGN_PARAGRAPH.LEFT)

            # run 级格式
            para_count = 0
            for run in para.runs:
                if target in run.text or run.text.strip() in target:
                    if "bold" in op:
                        run.bold = op["bold"]
                    if "italic" in op:
                        run.italic = op["italic"]
                    if "font_size" in op:
                        run.font.size = Pt(op["font_size"])
                    para_count += 1
            if para_count == 0:
                # 没有精确匹配 run，对所有 run 应用
                for run in para.runs:
                    if "bold" in op:
                        run.bold = op["bold"]
                    if "italic" in op:
                        run.italic = op["italic"]
                    if "font_size" in op:
                        run.font.size = Pt(op["font_size"])
                para_count = 1
            total_count += para_count
        return total_count

    @staticmethod
    def _apply_insert(doc: DocxDocument, after_text: str, content: str) -> int:
        if not content:
            return 0
        for i, para in enumerate(doc.paragraphs):
            if after_text and after_text in para.text:
                new_para = doc.add_paragraph(content)
                # 移动到目标段落之后
                para._element.addnext(new_para._element)
                return 1
        # 没找到目标，追加到末尾
        doc.add_paragraph(content)
        return 1

    @staticmethod
    def _parse_json_object(text: str) -> dict | None:
        text = text.strip()
        if text.startswith("```"):
            text = re.sub(r"^```\w*\n?", "", text)
            text = re.sub(r"\n?```$", "", text)
            text = text.strip()
        try:
            result = json.loads(text)
            return result if isinstance(result, dict) else None
        except json.JSONDecodeError:
            match = re.search(r"\{.*\}", text, flags=re.DOTALL)
            if match:
                try:
                    result = json.loads(match.group(0))
                    return result if isinstance(result, dict) else None
                except json.JSONDecodeError:
                    return None
            return None
