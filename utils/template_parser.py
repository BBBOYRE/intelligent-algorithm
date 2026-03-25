from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook

from core.document_parser import DocumentParser


class TemplateParser:
    """Parse template files into semantic markdown plus coordinate mappings."""

    def __init__(self) -> None:
        self.doc_parser = DocumentParser()

    def parse_template(self, file_path: str | Path) -> dict:
        path = Path(file_path)
        suffix = path.suffix.lower().lstrip(".")

        docling_result = self.doc_parser.parse(path)
        markdown = docling_result["markdown"]

        if suffix == "docx":
            coordinates = self._parse_word_coordinates(path)
        elif suffix in {"xlsx", "xls"}:
            coordinates = self._parse_excel_coordinates(path)
        else:
            raise ValueError(f"Unsupported template format: {suffix}")

        return {"markdown": markdown, "coordinates": coordinates}

    @staticmethod
    def _parse_word_coordinates(file_path: Path) -> list[dict]:
        doc = DocxDocument(file_path)
        table_infos: list[dict] = []
        for table_idx, table in enumerate(doc.tables):
            rows_data: list[list[str]] = []
            for row in table.rows:
                rows_data.append([cell.text.strip() for cell in row.cells])
            headers = rows_data[0] if rows_data else []
            table_infos.append(
                {
                    "table_index": table_idx,
                    "headers": headers,
                    "row_count": len(rows_data),
                    "structure": rows_data,
                }
            )
        return table_infos

    @staticmethod
    def _parse_excel_coordinates(file_path: Path) -> list[dict]:
        wb = load_workbook(file_path, read_only=True)
        sheets_info: list[dict] = []
        for ws in wb.worksheets:
            headers: dict[str, str] = {}
            for cell in ws[1]:
                if cell.value is not None and str(cell.value).strip():
                    headers[cell.coordinate] = str(cell.value).strip()
            sheets_info.append(
                {
                    "sheet_name": ws.title,
                    "headers": headers,
                    "header_row": 1,
                    "max_column": ws.max_column,
                }
            )
        wb.close()
        return sheets_info
