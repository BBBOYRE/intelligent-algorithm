"""格式转换工具：docx → txt/markdown, xlsx → csv/json"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path


def convert_file(source_path: str | Path, target_format: str) -> tuple[bytes, str]:
    """转换文件格式，返回 (bytes, suggested_filename)"""
    path = Path(source_path)
    suffix = path.suffix.lower()
    fmt = target_format.lower()

    if suffix == ".docx":
        if fmt == "txt":
            return _docx_to_txt(path), path.stem + ".txt"
        if fmt in ("md", "markdown"):
            return _docx_to_markdown(path), path.stem + ".md"

    if suffix in (".xlsx", ".xls"):
        if fmt == "csv":
            return _xlsx_to_csv(path), path.stem + ".csv"
        if fmt == "json":
            return _xlsx_to_json(path), path.stem + ".json"

    raise ValueError(f"不支持从 {suffix} 转换为 {fmt}")


def _docx_to_txt(path: Path) -> bytes:
    from docx import Document as DocxDocument
    doc = DocxDocument(path)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append("\t".join(cells))
    return "\n".join(lines).encode("utf-8")


def _docx_to_markdown(path: Path) -> bytes:
    from docx import Document as DocxDocument
    doc = DocxDocument(path)
    parts = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style and p.style.name.startswith("Heading"):
            level = int(p.style.name[-1]) if p.style.name[-1].isdigit() else 1
            parts.append("#" * level + " " + text)
        else:
            parts.append(text)
    for table in doc.tables:
        headers = [c.text.strip() for c in table.rows[0].cells]
        parts.append("| " + " | ".join(headers) + " |")
        parts.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            parts.append("| " + " | ".join(cells) + " |")
    return "\n\n".join(parts).encode("utf-8")


def _xlsx_to_csv(path: Path) -> bytes:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    buf = io.StringIO()
    writer = csv.writer(buf)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            writer.writerow([str(v) if v is not None else "" for v in row])
    wb.close()
    return buf.getvalue().encode("utf-8-sig")


def _xlsx_to_json(path: Path) -> bytes:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    result = {}
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[0])]
        data = []
        for row in rows[1:]:
            data.append({headers[i]: (str(v) if v is not None else "") for i, v in enumerate(row) if i < len(headers)})
        result[sheet.title] = data
    wb.close()
    return json.dumps(result, ensure_ascii=False, indent=2).encode("utf-8")
